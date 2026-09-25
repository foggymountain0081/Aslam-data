"""Write the same sheets to an Excel workbook (one tab per sheet)."""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from make_html import SHEETS, SUMMARY_COLS

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
HEAD = PatternFill("solid", fgColor="0B6E6E")
PROB = {"Very High": "0A6B2E", "High": "2F8F46", "Medium": "A86B00", "Low": "5F6B76", "Very Low": "B3261E"}


def _sheet(wb, title, rows, cols):
    ws = wb.create_sheet(title[:31])
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    if "Probability" in cols:
        pc = cols.index("Probability") + 1
        for row in ws.iter_rows(min_row=2, min_col=pc, max_col=pc):
            for c in row:
                if c.value in PROB:
                    c.font = Font(bold=True, color=PROB[c.value])
    for i, c in enumerate(cols, 1):
        width = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:300]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def build_xlsx(T, path):
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Area summary", T["summary"], SUMMARY_COLS)
    for key, label, cols in SHEETS:
        _sheet(wb, label.replace("–", "-"), T[key], cols)
    for area in T["stats"]["areas"]:  # area-wise: one tab per area with every list stacked
        ws = wb.create_sheet(f"Area - {area}"[:31])
        for key, label, cols in SHEETS:
            rows = [r for r in T[key] if r["Area"] == area]
            cols = [c for c in cols if c != "Area"]
            ws.append([f"{label} ({len(rows)})"])
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=12, color="0B6E6E")
            ws.append(cols)
            for c in ws[ws.max_row]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = HEAD
            for r in rows:
                ws.append([r.get(c, "") for c in cols])
            ws.append([])
        for i in range(1, 14):
            ws.column_dimensions[get_column_letter(i)].width = 24
    wb.save(path)


if __name__ == "__main__":
    from report_tables import make_tables
    build_xlsx(make_tables(), os.path.join(OUT, "Aslam_School_Report.xlsx"))
